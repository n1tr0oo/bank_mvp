"""
Конвертер SECURITY_ANALYSIS.md -> SECURITY_ANALYSIS.docx
Стиль: официальный, только чёрный шрифт, без заливки, без цветов.
"""
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import re


BLACK = RGBColor(0x00, 0x00, 0x00)
DARK_GRAY = RGBColor(0x44, 0x44, 0x44)


def set_table_borders(table):
    """Добавляет тонкие чёрные границы ко всей таблице."""
    tbl = table._tbl
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is None:
        tblPr = OxmlElement('w:tblPr')
        tbl.insert(0, tblPr)
    # Удаляем старые границы если есть
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    tblBorders = OxmlElement('w:tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        b = OxmlElement(f'w:{name}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '0')
        b.set(qn('w:color'), '000000')
        tblBorders.append(b)
    tblPr.append(tblBorders)


def remove_cell_shading(cell):
    """Удаляет заливку ячейки (белый фон)."""
    tc = cell._tc
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc.insert(0, tcPr)
    old = tcPr.find(qn('w:shd'))
    if old is not None:
        tcPr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'FFFFFF')
    tcPr.append(shd)


def add_code_block(doc, text):
    """Блок кода: Courier New 8pt, рамка, белый фон."""
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    run.font.name = 'Courier New'
    run.font.size = Pt(8)
    run.font.color.rgb = BLACK

    pPr = p._p.get_or_add_pPr()
    # Рамка вокруг блока
    pBdr = OxmlElement('w:pBdr')
    for side in ('top', 'left', 'bottom', 'right'):
        b = OxmlElement(f'w:{side}')
        b.set(qn('w:val'), 'single')
        b.set(qn('w:sz'), '4')
        b.set(qn('w:space'), '4')
        b.set(qn('w:color'), '888888')
        pBdr.append(b)
    pPr.append(pBdr)
    # Белый фон (без серой заливки)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), 'FFFFFF')
    pPr.append(shd)


def parse_inline(text):
    """Разбивает строку на части: normal, bold, code."""
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    result = []
    for part in parts:
        if part.startswith('`') and part.endswith('`'):
            result.append(('code', part[1:-1]))
        elif part.startswith('**') and part.endswith('**'):
            result.append(('bold', part[2:-2]))
        else:
            result.append(('normal', part))
    return result


def add_paragraph_with_inline(doc, text, style=None, bold=False):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    for kind, content in parse_inline(text):
        run = p.add_run(content)
        run.font.color.rgb = BLACK
        if kind == 'code':
            run.font.name = 'Courier New'
            run.font.size = Pt(9)
            run.font.color.rgb = BLACK
        elif kind == 'bold':
            run.bold = True
        if bold:
            run.bold = True
    return p


def set_run_black(run):
    run.font.color.rgb = BLACK


def configure_styles(doc):
    """Настраивает базовые стили документа — чёрный шрифт, без цветов."""
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = BLACK

    heading_sizes = {1: Pt(16), 2: Pt(14), 3: Pt(13), 4: Pt(12)}
    for level, pt in heading_sizes.items():
        name = f'Heading {level}'
        if name in doc.styles:
            s = doc.styles[name]
            s.font.name = 'Times New Roman'
            s.font.size = pt
            s.font.color.rgb = BLACK
            s.font.bold = True
            # Убираем тему/цвет через XML
            rPr = s.element.find('.//' + qn('w:rPr'))
            if rPr is not None:
                color_el = rPr.find(qn('w:color'))
                if color_el is not None:
                    rPr.remove(color_el)
                theme_color = rPr.find(qn('w:rFonts'))


def convert(md_path, docx_path):
    doc = Document()

    # Поля страницы
    section = doc.sections[0]
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.2)
    section.right_margin = Inches(0.8)

    configure_styles(doc)

    with open(md_path, encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    while i < len(lines):
        line = lines[i].rstrip('\n')

        # ---- Заголовки ----
        if line.startswith('#### '):
            h = doc.add_heading(line[5:], level=4)
            for run in h.runs:
                set_run_black(run)
            i += 1

        elif line.startswith('### '):
            h = doc.add_heading(line[4:], level=3)
            for run in h.runs:
                set_run_black(run)
            i += 1

        elif line.startswith('## '):
            h = doc.add_heading(line[3:], level=2)
            for run in h.runs:
                set_run_black(run)
            i += 1

        elif line.startswith('# '):
            h = doc.add_heading(line[2:], level=1)
            for run in h.runs:
                set_run_black(run)
            i += 1

        # ---- Блок кода ----
        elif line.startswith('```'):
            i += 1
            code_lines = []
            while i < len(lines) and not lines[i].rstrip('\n').startswith('```'):
                code_lines.append(lines[i].rstrip('\n'))
                i += 1
            add_code_block(doc, '\n'.join(code_lines))
            i += 1

        # ---- Таблица ----
        elif re.match(r'^\s*\|', line):
            table_lines = []
            while i < len(lines) and re.match(r'^\s*\|', lines[i]):
                table_lines.append(lines[i].rstrip('\n'))
                i += 1

            rows = []
            for tl in table_lines:
                # Пропускаем разделители |---|---|
                if re.match(r'^\s*\|[\s\-:|]+\|\s*$', tl):
                    continue
                cells = [c.strip() for c in tl.strip().strip('|').split('|')]
                rows.append(cells)

            if not rows:
                continue

            col_count = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=col_count)
            table.style = 'Table Grid'
            set_table_borders(table)

            for r_idx, row_data in enumerate(rows):
                for c_idx in range(col_count):
                    cell_text = row_data[c_idx] if c_idx < len(row_data) else ''
                    cell = table.cell(r_idx, c_idx)
                    remove_cell_shading(cell)
                    cell.text = ''
                    p = cell.paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for kind, content in parse_inline(cell_text):
                        run = p.add_run(content)
                        run.font.size = Pt(10)
                        run.font.name = 'Times New Roman'
                        run.font.color.rgb = BLACK
                        if r_idx == 0:
                            run.bold = True
                        if kind == 'bold':
                            run.bold = True
                        if kind == 'code':
                            run.font.name = 'Courier New'
                            run.font.size = Pt(9)

            doc.add_paragraph()

        # ---- Маркированный список ----
        elif line.startswith('- ') or line.startswith('* '):
            p = add_paragraph_with_inline(doc, line[2:], style='List Bullet')
            i += 1

        # ---- Нумерованный список ----
        elif re.match(r'^\d+\. ', line):
            text = re.sub(r'^\d+\. ', '', line)
            p = add_paragraph_with_inline(doc, text, style='List Number')
            i += 1

        # ---- Горизонтальная черта ----
        elif line.startswith('---'):
            p = doc.add_paragraph()
            pPr = p._p.get_or_add_pPr()
            pBdr = OxmlElement('w:pBdr')
            b = OxmlElement('w:bottom')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '6')
            b.set(qn('w:space'), '1')
            b.set(qn('w:color'), '000000')
            pBdr.append(b)
            pPr.append(pBdr)
            i += 1

        # ---- Пустая строка ----
        elif line.strip() == '':
            i += 1

        # ---- Обычный абзац ----
        else:
            add_paragraph_with_inline(doc, line)
            i += 1

    doc.save(docx_path)
    print(f"Сохранено: {docx_path}")


if __name__ == '__main__':
    convert('SECURITY_ANALYSIS.md', 'SECURITY_ANALYSIS.docx')
